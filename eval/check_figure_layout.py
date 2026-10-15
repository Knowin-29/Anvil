#!/usr/bin/env python3
"""Static layout audit for the editable figure PPTX.

Catches the two failure modes that are easy to miss by eye:
  * shapes escaping the slide canvas (LibreOffice silently clips them)
  * text boxes overlapping other text, which reads as garbled output in the PDF
and reports the greyscale luminance spread of the palette, since the figures must
survive black-and-white printing.

    python eval/check_figure_layout.py [--pptx docs/figures_editable_generated.pptx]

Exit code 1 if any hard problem (out-of-bounds, text-on-text overlap) is found.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

REPO = Path(__file__).resolve().parent.parent
EMU_IN = 914400.0
TOL = 0.012          # inches of slack before an overlap is called a problem


def rect(sh):
    return (sh.left / EMU_IN, sh.top / EMU_IN,
            (sh.left + sh.width) / EMU_IN, (sh.top + sh.height) / EMU_IN)


def has_text(sh) -> bool:
    return sh.has_text_frame and sh.text_frame.text.strip() != ""


def overlap(a, b) -> tuple[float, float]:
    return (min(a[2], b[2]) - max(a[0], b[0]), min(a[3], b[3]) - max(a[1], b[1]))


def audit(path: Path) -> int:
    prs = Presentation(str(path))
    W, H = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    print(f"{path.name}: {len(prs.slides._sldIdLst)} slides, canvas {W:.2f}x{H:.2f} in")
    problems = 0

    for idx, slide in enumerate(prs.slides, 1):
        shapes = [s for s in slide.shapes
                  if s.left is not None and s.width is not None]
        oob = []
        for s in shapes:
            x0, y0, x1, y1 = rect(s)
            # a connector may legitimately be zero-area; only flag canvas escapes
            if x0 < -TOL or y0 < -TOL or x1 > W + TOL or y1 > H + TOL:
                txt = (s.text_frame.text.strip().splitlines() or [""])[0][:44] \
                    if s.has_text_frame else "<no text>"
                oob.append((x0, y0, x1, y1, txt))

        texts = [s for s in shapes if has_text(s)]
        clashes = []
        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                ra, rb = rect(texts[i]), rect(texts[j])
                ox, oy = overlap(ra, rb)
                if ox > TOL and oy > TOL:
                    # a label sitting inside its own panel is fine only if the panel
                    # is empty; both carrying text means real collision risk
                    ta = (texts[i].text_frame.text.strip().splitlines() or [""])[0][:34]
                    tb = (texts[j].text_frame.text.strip().splitlines() or [""])[0][:34]
                    clashes.append((ox, oy, ta, tb))

        status = "OK" if not oob and not clashes else "PROBLEM"
        print(f"  slide {idx}: {len(shapes):3d} shapes, {len(texts):3d} with text  "
              f"[{status}]")
        for x0, y0, x1, y1, t in oob:
            print(f"      OUT OF BOUNDS ({x0:.2f},{y0:.2f})-({x1:.2f},{y1:.2f}) {t!r}")
        for ox, oy, ta, tb in clashes:
            print(f"      TEXT OVERLAP {ox:.3f}x{oy:.3f} in: {ta!r} <> {tb!r}")
        problems += len(oob) + len(clashes)

    print(f"\ntotal problems: {problems}")
    return 1 if problems else 0


def print_size_report(path: Path, textwidth_in: float = 5.5,
                      min_pt: float = 5.0) -> int:
    """Effective font size once LaTeX scales the figure to the ICLR text block.

    Placed at \\linewidth, the figure is scaled by textwidth / ink_width, and every
    glyph shrinks by the same factor. Anything landing under `min_pt` will not survive
    printing.
    """
    prs = Presentation(str(path))
    W = prs.slide_width / EMU_IN
    print(f"\neffective printed font size at ICLR \\linewidth ({textwidth_in}in):")
    worst = 0.0
    bad = 0
    for idx, slide in enumerate(prs.slides, 1):
        # measure real ink width, matching what the export crop produces
        xs = [(s.left / EMU_IN, (s.left + s.width) / EMU_IN)
              for s in slide.shapes if s.left is not None and s.width is not None]
        ink = max(b for _, b in xs) - min(a for a, _ in xs) if xs else W
        scale = textwidth_in / ink
        sizes = []
        for s in slide.shapes:
            if not (s.has_text_frame and s.text_frame.text.strip()):
                continue
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size is not None:
                        sizes.append(r.font.size.pt)
        if not sizes:
            continue
        smallest, largest = min(sizes) * scale, max(sizes) * scale
        flag = "OK" if smallest >= min_pt else "TOO SMALL"
        if smallest < min_pt:
            bad += 1
        worst = min(worst or smallest, smallest)
        print(f"  slide {idx}: ink {ink:.2f}in, scale {scale:.3f}  ->  "
              f"body {smallest:.1f}-{largest:.1f}pt  [{flag}]")
    print(f"  smallest printed text overall: {worst:.1f}pt "
          f"(floor {min_pt}pt)")
    return bad


def greyscale_report() -> None:
    import importlib
    m = importlib.import_module("make_figure_skeletons")
    names = ["INK", "MUT", "BOX", "TRK", "EDGE", "BLUE", "ORNG", "GRN", "RED", "WHITE"]
    print("\ngreyscale luminance (ITU-R BT.601), must stay separated for B/W print:")
    lum = {}
    for n in names:
        c = getattr(m, n)
        r, g, b = c[0], c[1], c[2]
        lum[n] = 0.299 * r + 0.587 * g + 0.114 * b
    for n, v in sorted(lum.items(), key=lambda kv: kv[1]):
        print(f"  {n:6s} {v:6.1f}  {'#' * int(v / 6)}")
    accents = {k: lum[k] for k in ("BLUE", "ORNG", "GRN", "RED")}
    order = sorted(accents.items(), key=lambda kv: kv[1])
    gaps = [(order[i + 1][1] - order[i][1]) for i in range(len(order) - 1)]
    print(f"  accent order: {' < '.join(k for k, _ in order)}")
    print(f"  min accent gap: {min(gaps):.1f} "
          f"({'OK' if min(gaps) >= 8 else 'TIGHT - rely on labels/length too'})")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx",
                    default=str(REPO / "docs" / "figures_editable_generated.pptx"))
    a = ap.parse_args()
    p = Path(a.pptx)
    if not p.exists():
        raise SystemExit(f"no such pptx: {p}")
    rc = audit(p)
    import sys
    sys.path.insert(0, str(REPO / "eval"))
    small = print_size_report(p)
    greyscale_report()
    return 1 if (rc or small) else 0


if __name__ == "__main__":
    raise SystemExit(main())
